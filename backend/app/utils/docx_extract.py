import io
import os
import zipfile
import xml.etree.ElementTree as ET
from typing import List, Union

def extract_text_from_docx(file_source: Union[str, bytes, io.BytesIO]) -> str:
    """
    Extracts raw text content from a DOCX file path, raw bytes, or stream object.
    Supports python-docx if installed, with robust standard-library zip/XML fallback.
    """
    # 1. Try python-docx if available
    try:
        import docx
        doc_source = io.BytesIO(file_source) if isinstance(file_source, bytes) else file_source
        doc = docx.Document(doc_source)
        paragraphs: List[str] = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        if paragraphs:
            return "\n".join(paragraphs)
    except Exception:
        pass

    # 2. Pure Python standard library fallback (zero extra dependencies required)
    try:
        zip_source = io.BytesIO(file_source) if isinstance(file_source, bytes) else file_source
        with zipfile.ZipFile(zip_source, "r") as z:
            if "word/document.xml" not in z.namelist():
                raise ValueError("Invalid DOCX format: word/document.xml missing")
            
            xml_data = z.read("word/document.xml")
            tree = ET.fromstring(xml_data)
            
            # w:p represents paragraphs, w:t represents text nodes
            w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            paragraphs = []
            for p in tree.iter(f"{{{w_ns}}}p"):
                texts = [node.text for node in p.iter(f"{{{w_ns}}}t") if node.text]
                if texts:
                    line = "".join(texts).strip()
                    if line:
                        paragraphs.append(line)
                        
            return "\n".join(paragraphs)
    except Exception as e:
        raise ValueError(f"Failed to read DOCX document: {str(e)}")
